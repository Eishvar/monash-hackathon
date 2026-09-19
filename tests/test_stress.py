"""Synthetic stress tests: generate SI/BL documents in several formats and label vocabularies, then check the
properties the product promises: equivalent formatting never alarms, a single real change flags exactly that field,
blank values escalate, wrong document types escalate."""
import io

import pytest

from sdoc.config import FIELDS
from sdoc.decide import decide
from sdoc.normalize import norm_count, norm_party, norm_weight
from sdoc.parse import read_document

BASE = {
    "shipper": "ACME TRADING PTE LIMITED\n1 MAIN ROAD; SINGAPORE 048624",
    "consignee": "BUYER GMBH\nOPERNRING 3; 1010 VIENNA, AUSTRIA",
    "notify_party": "NOTIFY LOGISTICS CO., LTD",
    "port_of_loading": "SINGAPORE",
    "port_of_discharge": "HAMBURG, GERMANY (DEHAM)",
    "container_count": "6 x 40'HC",
    "gross_weight_kg": "131,058 KG",
}
CHANGED = {
    "shipper": "OTHER TRADING PTE LTD\n1 MAIN ROAD; SINGAPORE 048624",
    "consignee": "DIFFERENT BUYER GMBH\nOPERNRING 3; 1010 VIENNA, AUSTRIA",
    "notify_party": "ANOTHER NOTIFY LTD",
    "port_of_loading": "PORT KLANG",
    "port_of_discharge": "ROTTERDAM, NETHERLANDS",
    "container_count": "7 x 40'HC",
    "gross_weight_kg": "131,059 KG",
}
LABELS = {
    "A": dict(zip(FIELDS, ["Shipper/Exporter", "To the Order of", "Notify", "Load Port", "Discharge Port",
                           "Total Containers", "Gross Wt (kgs)"])),
    "B": dict(zip(FIELDS, ["SHIPPER", "Consignee (Non-Negotiable)", "NOTIFY PARTY", "Port of Loading (POL)",
                           "Port of Discharge (POD)", "No. of Containers or Packages", "GROSS WEIGHT"])),
    "C": dict(zip(FIELDS, ["Shipper (Principal or Seller) (发货人)", "Consignee (收货人)", "Notify Party/Intermediate Consignee",
                           "POL", "POD", "Container Count", "Gross Weight毛重(KGS)"])),
}
NOISE = [("Vessel", "MV TEST V.1"), ("NET WEIGHT", "1,000 KG"), ("Booking No.", "BK123"), ("HS CODE", "4802")]


def build_txt(title, labels, values) -> bytes:
    lines = [title, "=" * 40, ""]
    for f in FIELDS:
        first, *rest = values[f].split("\n")
        lines.append(f"{labels[f]}: {first}")
        lines += [f"  {r}" for r in rest]
    lines += [f"{k}: {v}" for k, v in NOISE]
    return "\n".join(lines).encode("utf-8")


def build_docx(title, labels, values) -> bytes:
    import docx

    d = docx.Document()
    d.add_paragraph(title)
    t = d.add_table(rows=0, cols=2)
    for f in FIELDS:
        row = t.add_row().cells
        row[0].text, row[1].text = labels[f], values[f]
    for k, v in NOISE:
        row = t.add_row().cells
        row[0].text, row[1].text = k, v
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


def build_xlsx(title, labels, values) -> bytes:
    import openpyxl

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["SOME COMPANY", None])
    ws.append([title, "REF123"])
    for f in FIELDS:
        v = values[f]
        if f == "gross_weight_kg" and v.replace(",", "").replace(" KG", "").isdigit():
            v = int(v.replace(",", "").replace(" KG", ""))  # numeric cell, like the real files
        ws.append([labels[f], v.replace("\n", " | ") if isinstance(v, str) else v])
    for k, v in NOISE:
        ws.append([k, v])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


BUILDERS = {"txt": build_txt, "docx": build_docx, "xlsx": build_xlsx}
FORMATS = list(BUILDERS)


def pair(fmt, si_labels="A", bl_labels="B", bl_values=None, bl_title="BILL OF LADING (DRAFT)"):
    si = read_document(f"x_SI.{fmt}", BUILDERS[fmt]("SHIPPING INSTRUCTION", LABELS[si_labels], BASE))
    bl = read_document(f"x_BL.{fmt}", BUILDERS[fmt](bl_title, LABELS[bl_labels], bl_values or BASE))
    return si, bl


@pytest.mark.parametrize("fmt", FORMATS)
@pytest.mark.parametrize("si_l,bl_l", [("A", "B"), ("B", "C"), ("C", "A")])
def test_identical_content_different_labels_is_ok(fmt, si_l, bl_l):
    si, bl = pair(fmt, si_l, bl_l)
    assert si.kind == "SI" and bl.kind == "BL"
    assert all(v for v in si.fields.values()), si.fields  # all 7 found despite noise labels (NET WEIGHT, ...)
    d = decide(si, bl)
    assert (d.status, d.defect_fields) == ("OK", [])


@pytest.mark.parametrize("fmt", FORMATS)
@pytest.mark.parametrize("field", FIELDS)
def test_single_real_change_flags_exactly_that_field(fmt, field):
    si, bl = pair(fmt, bl_values={**BASE, field: CHANGED[field]})
    d = decide(si, bl)
    assert (d.status, d.defect_fields) == ("MISMATCH", [field])


@pytest.mark.parametrize("fmt", FORMATS)
def test_formatting_differences_do_not_alarm(fmt):
    fmt_only = {
        **BASE,
        "shipper": "Acme Trading Pte Ltd\n1 Main Road, Singapore 048624",  # case, punctuation, LIMITED vs LTD
        "port_of_discharge": "HAMBURG",  # no country / LOCODE
        "container_count": "Six (6) x 40'HC",
        "gross_weight_kg": "131,058.0 KGS",
    }
    d = decide(*pair(fmt, bl_values=fmt_only))
    assert (d.status, d.defect_fields) == ("OK", [])


@pytest.mark.parametrize("fmt", FORMATS)
@pytest.mark.parametrize("blank", ["", "N/A", "TBA", "________", "-"])
def test_blank_or_placeholder_escalates_as_missing_value(fmt, blank):
    si, bl = pair(fmt, bl_values={**BASE, "consignee": blank})
    d = decide(si, bl)
    assert (d.status, d.review_reason, d.defect_fields) == ("NEEDS_REVIEW", "missing_value", [])


@pytest.mark.parametrize("fmt", FORMATS)
@pytest.mark.parametrize("title", ["COMMERCIAL INVOICE", "PACKING LIST", "CERTIFICATE OF ORIGIN"])
def test_wrong_document_type_escalates(fmt, title):
    si, bl = pair(fmt, bl_title=title)
    assert decide(si, bl).review_reason == "wrong_doc_type"


def test_unit_and_word_normalisation():
    assert norm_weight("22 MT") == norm_weight("22,000 KG") == 22000
    assert norm_weight("22.5 tonnes") == 22500
    assert norm_weight("1,234.5 KGS") == 1234.5
    assert norm_count("Six (6) x 40'HC") == norm_count("six containers") == 6
    assert norm_party("ACME LIMITED") == norm_party("Acme Ltd.")
    assert norm_party("Smith & Sons Company") == norm_party("SMITH AND SONS CO")
    assert norm_party("ACME TRADING") != norm_party("ACME TRADERS")
