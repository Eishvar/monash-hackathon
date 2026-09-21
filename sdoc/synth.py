"""Synthetic SI / draft-BL documents whose correct verdict is known by construction.

Used by the stress tests (fixed BASE/CHANGED pair) and by `scripts/benchmark.py`, which generates thousands of
random pairs to measure the comparison engine on data it was never tuned on: equivalent formatting must stay OK,
planted defects must be found exactly, blanks / wrong documents / missing / unreadable files must escalate."""
import io
import random

from sdoc.config import FIELDS
from sdoc.samples import text_pdf

BASE = {
    "shipper": "ACME TRADING PTE LIMITED\n1 MAIN ROAD; SINGAPORE 048624",
    "consignee": "BUYER GMBH\nOPERNRING 3; 1010 VIENNA, AUSTRIA",
    "notify_party": "NOTIFY LOGISTICS CO., LTD",
    "port_of_loading": "SINGAPORE",
    "port_of_discharge": "HAMBURG, GERMANY (DEHAM)",
    "container_count": "6 x 40'HC",
    "gross_weight_kg": "131,058 KG",
}
CHANGED = {
    "shipper": "OTHER TRADING PTE LTD\n1 MAIN ROAD; SINGAPORE 048624",
    "consignee": "DIFFERENT BUYER GMBH\nOPERNRING 3; 1010 VIENNA, AUSTRIA",
    "notify_party": "ANOTHER NOTIFY LTD",
    "port_of_loading": "PORT KLANG",
    "port_of_discharge": "ROTTERDAM, NETHERLANDS",
    "container_count": "7 x 40'HC",
    "gross_weight_kg": "131,059 KG",
}
LABELS = {
    "A": dict(zip(FIELDS, ["Shipper/Exporter", "To the Order of", "Notify", "Load Port", "Discharge Port",
                           "Total Containers", "Gross Wt (kgs)"])),
    "B": dict(zip(FIELDS, ["SHIPPER", "Consignee (Non-Negotiable)", "NOTIFY PARTY", "Port of Loading (POL)",
                           "Port of Discharge (POD)", "No. of Containers or Packages", "GROSS WEIGHT"])),
    "C": dict(zip(FIELDS, ["Shipper (Principal or Seller) (发货人)", "Consignee (收货人)", "Notify Party/Intermediate Consignee",
                           "POL", "POD", "Container Count", "Gross Weight毛重(KGS)"])),
}
NOISE = [("Vessel", "MV TEST V.1"), ("NET WEIGHT", "1,000 KG"), ("Booking No.", "BK123"), ("HS CODE", "4802")]


def build_txt(title, labels, values) -> bytes:
    return "\n".join(_lines(title, labels, values)).encode("utf-8")


def _lines(title, labels, values) -> list[str]:
    lines = [title, "=" * 40, ""]
    for f in FIELDS:
        first, *rest = values[f].split("\n")
        lines.append(f"{labels[f]}: {first}")
        lines += [f"  {r}" for r in rest]
    return lines + [f"{k}: {v}" for k, v in NOISE]


def build_docx(title, labels, values) -> bytes:
    import docx

    d = docx.Document()
    d.add_paragraph(title)
    t = d.add_table(rows=0, cols=2)
    for f in FIELDS:
        row = t.add_row().cells
        row[0].text, row[1].text = labels[f], values[f]
    for k, v in NOISE:
        row = t.add_row().cells
        row[0].text, row[1].text = k, v
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


def build_xlsx(title, labels, values) -> bytes:
    import openpyxl

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["SOME COMPANY", None])
    ws.append([title, "REF123"])
    for f in FIELDS:
        v = values[f]
        if f == "gross_weight_kg" and v.replace(",", "").replace(" KG", "").isdigit():
            v = int(v.replace(",", "").replace(" KG", ""))  # numeric cell, like the real files
        ws.append([labels[f], v.replace("\n", " | ") if isinstance(v, str) else v])
    for k, v in NOISE:
        ws.append([k, v])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def build_pdf(title, labels, values) -> bytes:
    """Text-layer PDF with the same line layout as the .txt builder (blank lines dropped, like the sample PDFs)."""
    lines = [ln.encode("latin-1", errors="ignore").decode("latin-1") for ln in _lines(title, labels, values)]  # Helvetica has no CJK
    return text_pdf([ln for ln in lines if ln.strip() and set(ln.strip()) != {"="}])


BUILDERS = {"txt": build_txt, "docx": build_docx, "xlsx": build_xlsx, "pdf": build_pdf}
FORMATS = list(BUILDERS)

# -- random cases with known answers ---------------------------------------------------------------------------------
# (name, one-word-different name, address line 1, address line 2)
PARTIES = [
    ("ACME TRADING PTE LIMITED", "ACME TRADERS PTE LIMITED", "1 MAIN ROAD", "SINGAPORE 048624"),
    ("BLUE HARBOUR EXPORTS CO., LTD", "BLUE HARBOUR IMPORTS CO., LTD", "88 DOCK STREET", "SHANGHAI 200000, CHINA"),
    ("GREENFIELD PAPER & PULP LIMITED", "GREENFIELD PAPER & PACKAGING LIMITED", "12 MILL LANE", "JAKARTA 10110, INDONESIA"),
    ("NORTHGATE PAPER MILLS PTE LTD", "NORTHGATE PAPER MARKETS PTE LTD", "5 TUAS AVENUE", "SINGAPORE 638000"),
    ("ORIENT LOGISTICS COMPANY", "ORIENT FREIGHT COMPANY", "77 HARBOUR ROAD", "HONG KONG"),
    ("BUYER GMBH", "SELLER GMBH", "OPERNRING 3", "1010 VIENNA, AUSTRIA"),
    ("ALDERMERE IMPORTS FZE", "ALDERMERE EXPORTS FZE", "PLOT 14 JAFZA", "DUBAI, UAE"),
    ("SUNRISE PACKAGING & SUPPLIES INC", "SUNRISE PACKAGING & SERVICES INC", "300 BAY STREET", "TORONTO ON, CANADA"),
]
PORTS = [("SINGAPORE", "SINGAPORE", "SGSIN"), ("HAMBURG", "GERMANY", "DEHAM"), ("ROTTERDAM", "NETHERLANDS", "NLRTM"),
         ("PORT KLANG", "MALAYSIA", "MYPKG"), ("SHANGHAI", "CHINA", "CNSHA"), ("BUSAN", "SOUTH KOREA", "KRPUS"),
         ("JEBEL ALI", "UAE", "AEJEA"), ("NHAVA SHEVA", "INDIA", "INNSA")]
CONTAINER_TYPES = ["40'HC", "20'GP", "40'GP"]
NUMBER_WORDS = "ONE TWO THREE FOUR FIVE SIX SEVEN EIGHT NINE TEN ELEVEN TWELVE".split()
BLANKS = ["", "N/A", "TBA", "________", "-"]
WRONG_TITLES = ["COMMERCIAL INVOICE", "PACKING LIST"]
SI_TITLE, BL_TITLE = "SHIPPING INSTRUCTION", "BILL OF LADING (DRAFT)"
CASE_MIX = [("clean", 0.40), ("defect", 0.35), ("near_miss", 0.05), ("blank", 0.07), ("wrong_doc", 0.05),
            ("missing", 0.04), ("unreadable", 0.04)]


def _party_value(p: tuple, name_idx: int = 0) -> str:
    return f"{p[name_idx]}\n{p[2]}; {p[3]}"


def random_values(rng: random.Random) -> dict:
    port = lambda p: f"{p[0]}, {p[1]} ({p[2]})"  # noqa: E731
    pol, pod = rng.sample(PORTS, 2)
    shipper, consignee, notify = rng.sample(PARTIES, 3)
    n = rng.randint(1, 12)
    return {
        "shipper": _party_value(shipper), "consignee": _party_value(consignee), "notify_party": _party_value(notify),
        "port_of_loading": port(pol), "port_of_discharge": port(pod),
        "container_count": f"{n} x {rng.choice(CONTAINER_TYPES)}",
        "gross_weight_kg": f"{rng.randrange(5000, 250000, 10):,} KG",
    }


def _weight(value: str) -> int:
    return int(value.split()[0].replace(",", ""))


def equivalent_variant(field: str, value: str, rng: random.Random) -> str:
    """Rewrites `value` in a way the spec treats as equal (formatting only)."""
    if field in ("shipper", "consignee", "notify_party"):
        options = [lambda v: v.title().replace("Gmbh", "GmbH"),
                   lambda v: v.replace("LIMITED", "LTD") if "LIMITED" in v else v.replace(" LTD", " LIMITED"),
                   lambda v: v.replace("&", "AND") if "&" in v else v.replace(" AND ", " & "),
                   lambda v: v.replace("; ", ", ").replace("\n", " | "),
                   lambda v: v.replace(" COMPANY", " CO.") if " COMPANY" in v else v.replace(" CO.", " COMPANY")]
        return rng.choice(options)(value)
    if field in ("port_of_loading", "port_of_discharge"):
        name, _, rest = value.partition(",")
        country = rest.partition("(")[0].strip()
        return rng.choice([name, f"{name}, {country}", f"{name} ({rest.rpartition('(')[2].rstrip(')')})", name.title()])
    if field == "container_count":
        n, _, kind = value.partition(" x ")
        return rng.choice([f"{NUMBER_WORDS[int(n) - 1].title()} ({n}) x {kind}", f"{n} X {kind}", f"{n} x {kind.lower()}"])
    kg = _weight(value)
    return rng.choice([f"{kg} KGS", f"{kg:,}.00 KG", f"{kg / 1000:g} MT" if kg % 10 == 0 else f"{kg} KG"])


def real_change(field: str, value: str, rng: random.Random) -> str:
    """A different value: another party / port, a container count off by 1-3, a weight off by 1-5,000 kg."""
    if field in ("shipper", "consignee", "notify_party"):
        others = [p for p in PARTIES if _party_value(p) != value]
        return _party_value(rng.choice(others))
    if field in ("port_of_loading", "port_of_discharge"):
        current = value.partition(",")[0]
        p = rng.choice([p for p in PORTS if p[0] != current])
        return f"{p[0]}, {p[1]} ({p[2]})"
    if field == "container_count":
        n, _, kind = value.partition(" x ")
        new = max(1, int(n) + rng.choice([-3, -2, -1, 1, 2, 3]))
        return f"{new if new != int(n) else new + 1} x {kind}"
    return f"{max(1, _weight(value) + rng.choice([-1, 1]) * rng.randint(1, 5000)):,} KG"


def _near_miss(field: str, value: str) -> str:
    """One word of the party name changed (ACME TRADING -> ACME TRADERS)."""
    for p in PARTIES:
        if _party_value(p) == value:
            return _party_value(p, 1)
    raise KeyError(value)


def _corrupt_pdf(rng: random.Random) -> bytes:
    return b"%PDF-1.4\n" + bytes(rng.randrange(256) for _ in range(300))


def make_case(i: int, rng: random.Random) -> dict:
    """One benchmark case: attachment files plus the verdict a correct engine must produce."""
    from sdoc.simulate import _scan_pdf  # local: pulls in Pillow only when a scan is built

    r, acc, case_type = rng.random(), 0.0, CASE_MIX[-1][0]
    for name, share in CASE_MIX:
        acc += share
        if r < acc:
            case_type = name
            break
    si_vals = random_values(rng)
    bl_vals = dict(si_vals)
    fmt_si, fmt_bl = rng.choice(FORMATS), rng.choice(FORMATS)
    lab_si, lab_bl = rng.choice("ABC"), rng.choice("ABC")
    bl_title, expected = BL_TITLE, {"status": "OK", "review_reason": None, "defect_fields": []}

    def vary(fields):
        for f in fields:
            bl_vals[f] = equivalent_variant(f, bl_vals[f], rng)

    if case_type in ("clean", "defect", "near_miss"):
        changed: list[str] = []
        if case_type == "defect":
            changed = rng.sample(FIELDS, rng.randint(1, 3))
            for f in changed:
                bl_vals[f] = real_change(f, si_vals[f], rng)
        elif case_type == "near_miss":
            changed = [rng.choice(["shipper", "consignee", "notify_party"])]
            bl_vals[changed[0]] = _near_miss(changed[0], si_vals[changed[0]])
        pool = [f for f in FIELDS if f not in changed]
        vary(rng.sample(pool, min(len(pool), rng.randint(1, 4) if case_type == "clean" else rng.randint(0, 2))))
        if changed:
            expected = {"status": "MISMATCH", "review_reason": None, "defect_fields": sorted(changed)}
    elif case_type == "blank":
        bl_vals[rng.choice(FIELDS)] = rng.choice(BLANKS)
        expected = {"status": "NEEDS_REVIEW", "review_reason": "missing_value", "defect_fields": []}
    elif case_type == "wrong_doc":
        bl_title = rng.choice(WRONG_TITLES)
        expected = {"status": "NEEDS_REVIEW", "review_reason": "wrong_doc_type", "defect_fields": []}
    elif case_type == "missing":
        expected = {"status": "NEEDS_REVIEW", "review_reason": "missing_attachment", "defect_fields": []}
    else:
        expected = {"status": "NEEDS_REVIEW", "review_reason": "unreadable", "defect_fields": []}

    si_name, bl_name = f"attachments/case_{i:04d}_SI.{fmt_si}", f"attachments/case_{i:04d}_BL.{fmt_bl}"
    files = {si_name: BUILDERS[fmt_si](SI_TITLE, LABELS[lab_si], si_vals)}
    if case_type == "unreadable":
        fmt_bl = rng.choice(["corrupt", "scan"])
        bl_name = f"attachments/case_{i:04d}_BL.pdf"
        files[bl_name] = _corrupt_pdf(rng) if fmt_bl == "corrupt" else _scan_pdf("\n".join(_lines(BL_TITLE, LABELS[lab_bl], bl_vals)))
    elif case_type != "missing":
        files[bl_name] = BUILDERS[fmt_bl](bl_title, LABELS[lab_bl], bl_vals)
    return {"id": f"case_{i:04d}", "case_type": case_type, "format_si": fmt_si,
            "format_bl": "-" if case_type == "missing" else fmt_bl, "files": files,
            "attachments": [n for n in (si_name, bl_name) if n in files], "expected": expected}
