"""Read txt/docx/xlsx/pdf attachments into a `Doc` (kind + the 7 raw field values).

Every reader produces (label, value) pairs; `_doc_from_pairs` maps labels to fields by meaning.
Unreadable input (corrupt file, image-only scan) never raises: it returns `Doc(unreadable=True)`.
"""
import io
import re
from dataclasses import dataclass, field

from sdoc.config import FIELDS
from sdoc.parse.labels import NOISE_PREFIX, field_for_label, match_line_label, strip_cjk

MIN_PDF_TEXT_CHARS = 20
_BLANK = re.compile(
    r"^[\s_\-–—./]*$|^n\s*/?\s*a$|^nil$|^tb[acd]$|^to be (advised|confirmed|determined)$|_{3,}", re.IGNORECASE
)
_TABLE_ROW = re.compile(r"^[A-Z]{4}\d{7}\b.*?(\d[\d,]*(?:\.\d+)?)\s*$")


@dataclass
class Doc:
    kind: str  # SI | BL | INVOICE | PACKING_LIST | COO | UNKNOWN
    fields: dict[str, str | None] = field(default_factory=lambda: dict.fromkeys(FIELDS))
    unreadable: bool = False
    error: str | None = None
    text: str = ""  # plain-text rendering of the document (input for LLM extraction)
    scanned: bool = False  # PDF with no text layer: candidate for the vision LLM


def is_blank(value: str | None) -> bool:
    return value is None or bool(_BLANK.match(value.strip()))


def detect_kind(head: str) -> str:
    h = head.lower()
    if "commercial invoice" in h:
        return "INVOICE"
    if "packing list" in h:
        return "PACKING_LIST"
    if "certificate of origin" in h:
        return "COO"
    if "instruction" in h:
        return "SI"
    if "bill of lading" in h:
        return "BL"
    return "UNKNOWN"


def _doc_from_pairs(head: str, pairs: list[tuple[str, str]], text: str = "") -> Doc:
    doc = Doc(kind=detect_kind(head), text=text)
    for label, value in pairs:
        f = field_for_label(label)
        if f and doc.fields[f] is None:
            doc.fields[f] = None if is_blank(value) else value.strip()
    return doc


def read_txt(data: bytes) -> Doc:
    lines = data.decode("utf-8", errors="replace").splitlines()
    head = " ".join(ln for ln in lines[:6] if ln.strip())
    pairs: list[list[str]] = []
    for ln in lines:
        if not ln.strip() or set(ln.strip()) == {"="}:
            continue
        if ln[0].isspace() and pairs:
            pairs[-1][1] += "\n" + ln.strip()
        elif ":" in ln:
            label, _, value = ln.partition(":")
            pairs.append([label, value.strip()])
    return _doc_from_pairs(head, [(a, b) for a, b in pairs], text="\n".join(lines))


def read_docx(data: bytes) -> Doc:
    import docx

    d = docx.Document(io.BytesIO(data))
    paras = [p.text.strip() for p in d.paragraphs if p.text.strip()]
    pairs = [(row.cells[0].text, row.cells[1].text) for t in d.tables for row in t.rows if len(row.cells) >= 2]
    for p in paras:
        if ":" in p:
            label, _, value = p.partition(":")
            pairs.append((label, value))
    text = "\n".join(paras + [f"{a.strip()}: {b.strip()}" for a, b in pairs if a.strip()])
    return _doc_from_pairs(" ".join(paras[:6]), pairs, text=text)


def _cell(v) -> str:
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    return str(v).strip()


def read_xlsx(data: bytes) -> Doc:
    import openpyxl

    ws = openpyxl.load_workbook(io.BytesIO(data), data_only=True).worksheets[0]
    rows = [tuple(_cell(c) for c in row[:2]) for row in ws.iter_rows(values_only=True)]
    rows = [r for r in rows if any(r)]
    head = " ".join(" ".join(r) for r in rows[:5])
    text = "\n".join(f"{r[0]}: {r[1]}" if len(r) > 1 and r[1] else r[0] for r in rows)
    return _doc_from_pairs(head, [(r[0], r[1] if len(r) > 1 else "") for r in rows], text=text)


def _pdf_lines(data: bytes) -> list[str]:
    import pdfplumber

    # Words in content-stream order (not x-sorted characters): a wrapped label column can overlap the value
    # column, and extract_text() would interleave their characters.
    lines: list[str] = []
    with pdfplumber.open(io.BytesIO(data)) as pdf:
        for page in pdf.pages:
            top, words = None, []
            for w in page.extract_words(use_text_flow=True):
                if top is not None and abs(w["top"] - top) > 3:
                    lines.append(" ".join(words))
                    words = []
                top = w["top"]
                words.append(w["text"])
            if words:
                lines.append(" ".join(words))
    return [strip_cjk(ln).strip() for ln in lines if ln.strip()]


def read_pdf(data: bytes) -> Doc:
    try:
        lines = _pdf_lines(data)
    except Exception as exc:  # corrupt / truncated stream
        return Doc(kind="UNKNOWN", unreadable=True, error=f"{type(exc).__name__}: {exc}")
    if sum(len(ln) for ln in lines) < MIN_PDF_TEXT_CHARS:
        return Doc(kind="UNKNOWN", unreadable=True, scanned=True, error="no text layer (image-only scan)")

    pairs: list[list[str]] = []
    current = False
    rows = 0
    row_weight = 0.0
    for ln in lines[1:]:
        m = match_line_label(ln)
        if m:
            pairs.append([m[0], m[1]])
            current = True
        elif (t := _TABLE_ROW.match(ln)) is not None:
            rows += 1
            row_weight += float(t.group(1).replace(",", ""))
            current = False
        elif NOISE_PREFIX.match(ln):
            current = False
        elif current:
            pairs[-1][1] += "\n" + ln

    doc = _doc_from_pairs(" ".join(lines[:4]), [], text="\n".join(lines))
    for f, value in pairs:
        if doc.fields[f] is None:
            doc.fields[f] = None if is_blank(value) else value.strip()
    if doc.fields["container_count"] is None and rows:
        doc.fields["container_count"] = str(rows)
    if doc.fields["gross_weight_kg"] is None and rows:
        doc.fields["gross_weight_kg"] = f"{row_weight:g}"
    return doc


def read_document(filename: str, data: bytes) -> Doc:
    ext = filename.rsplit(".", 1)[-1].lower()
    try:
        if ext == "txt":
            return read_txt(data)
        if ext == "docx":
            return read_docx(data)
        if ext == "xlsx":
            return read_xlsx(data)
        if ext == "pdf":
            return read_pdf(data)
    except Exception as exc:
        return Doc(kind="UNKNOWN", unreadable=True, error=f"{type(exc).__name__}: {exc}")
    return Doc(kind="UNKNOWN", unreadable=True, error=f"unsupported file type .{ext}")
